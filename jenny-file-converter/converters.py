import os, subprocess, shutil, csv
from pathlib import Path
from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter, PdfMerger
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
import pdfplumber

ALLOWED = {
    "pdf": {"docx","xlsx","pptx"},
    "docx": {"pdf"}, "xlsx": {"pdf","csv"}, "pptx": {"pdf"},
    "txt": {"pdf","docx"}, "jpg": {"png","pdf"}, "jpeg": {"png","pdf"},
    "png": {"jpg","jpeg","pdf"}, "webp": {"png","jpg","jpeg"},
    "csv": {"xlsx"}, "wav": {"mp3"}, "mp3": {"wav"},
    "mp4": {"mp3","wav","m4a","aac","flac"}, "mkv": {"mp3","wav","m4a","aac","flac"},
    "mov": {"mp3","wav","m4a","aac","flac"}, "avi": {"mp3","wav","m4a","aac","flac"},
    "webm": {"mp3","wav","m4a","aac","flac"},
}
def run(cmd):
    p = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=600)
    if p.returncode:
        raise RuntimeError((p.stderr or p.stdout or "Conversion failed")[-1200:])
def libreoffice(src, outdir, target):
    exe = shutil.which("libreoffice") or shutil.which("soffice")
    if not exe: raise RuntimeError("LibreOffice is not installed on this server.")
    run([exe, "--headless", "--convert-to", target, "--outdir", str(outdir), str(src)])
    out = outdir / (src.stem + "." + target)
    if not out.exists(): raise RuntimeError("LibreOffice could not convert this document.")
    return out
def pdf_to_docx(src, outdir):
    try:
        from pdf2docx import Converter
        out = outdir / (src.stem + ".docx")
        cv = Converter(str(src)); cv.convert(str(out)); cv.close()
        return out
    except Exception as e:
        raise RuntimeError("PDF-to-Word conversion failed. Scanned PDFs may need OCR. " + str(e)[:250])
def pdf_to_xlsx(src, outdir):
    out = outdir / (src.stem + ".xlsx"); wb = Workbook(); ws = wb.active
    found = False
    with pdfplumber.open(src) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables() or []
            for table in tables:
                found = True
                ws.append([f"Page {page_no}"])
                for row in table:
                    ws.append([cell or "" for cell in row])
                ws.append([])
        if not found:
            ws.append(["No tables detected. Extracted text follows; layout may be approximate."])
            for page in pdf.pages:
                for line in (page.extract_text() or "").splitlines(): ws.append([line])
    wb.save(out); return out
def pdf_to_pptx(src, outdir):
    import fitz
    pdf = fitz.open(src); prs = Presentation()
    prs.slide_width = Inches(10); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for page in pdf:
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5,1.5), alpha=False)
        image = outdir / f"page-{len(prs.slides)+1}.png"; pix.save(image)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
    out = outdir / (src.stem + ".pptx"); prs.save(out); return out
def image_convert(src, fmt, outdir):
    fmt = "JPEG" if fmt in ("jpg","jpeg") else fmt.upper()
    out = outdir / (src.stem + (".jpg" if fmt == "JPEG" else "."+fmt.lower()))
    with Image.open(src) as im:
        im = ImageOps.exif_transpose(im)
        if fmt == "JPEG" and im.mode in ("RGBA","LA","P"):
            rgba = im.convert("RGBA"); bg = Image.new("RGB", rgba.size, "white"); bg.paste(rgba, mask=rgba.getchannel("A")); im = bg
        elif fmt == "JPEG": im = im.convert("RGB")
        im.save(out, format=fmt)
    return out
def text_to_docx(src, outdir):
    d = Document()
    for line in src.read_text(encoding="utf-8", errors="replace").splitlines(): d.add_paragraph(line)
    out=outdir/(src.stem+".docx"); d.save(out); return out
def text_to_pdf(src, outdir):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    out=outdir/(src.stem+".pdf"); c=canvas.Canvas(str(out), pagesize=letter)
    w,h=letter; y=h-50
    for line in src.read_text(encoding="utf-8", errors="replace").splitlines():
        # Basic text output; long lines are clipped and formatting is not preserved.
        c.drawString(40,y,line[:120]); y-=15
        if y<45: c.showPage(); y=h-50
    c.save(); return out
def csv_to_xlsx(src, outdir):
    wb=Workbook(); ws=wb.active
    with open(src, newline="", encoding="utf-8-sig", errors="replace") as f:
        for row in csv.reader(f): ws.append(row)
    out=outdir/(src.stem+".xlsx"); wb.save(out); return out
def xlsx_to_csv(src, outdir):
    wb=load_workbook(src, read_only=True, data_only=True); ws=wb.active
    out=outdir/(src.stem+".csv")
    with open(out,"w",newline="",encoding="utf-8") as f:
        writer=csv.writer(f)
        for row in ws.iter_rows(values_only=True): writer.writerow([v if v is not None else "" for v in row])
    return out
def audio_convert(src, fmt, outdir, bitrate):
    exe=shutil.which("ffmpeg")
    if not exe: raise RuntimeError("FFmpeg is not installed on this server.")
    ext=fmt; out=outdir/(src.stem+"."+ext)
    cmd=[exe,"-y","-i",str(src),"-vn"]
    if fmt=="mp3": cmd += ["-codec:a","libmp3lame","-b:a", bitrate if bitrate in ("128k","192k","320k") else "192k"]
    elif fmt=="wav": cmd += ["-codec:a","pcm_s16le"]
    elif fmt=="m4a": cmd += ["-codec:a","aac","-b:a",bitrate if bitrate in ("128k","192k","320k") else "192k"]
    elif fmt=="aac": cmd += ["-codec:a","aac","-b:a",bitrate if bitrate in ("128k","192k","320k") else "192k"]
    elif fmt=="flac": cmd += ["-codec:a","flac"]
    cmd += [str(out)]; run(cmd); return out
def convert_file(paths, fmt, outdir, bitrate="192k"):
    if len(paths)!=1: raise RuntimeError("This conversion accepts one file at a time.")
    src=paths[0]; ext=src.suffix.lower().lstrip(".")
    if fmt not in ALLOWED.get(ext,set()): raise RuntimeError(f"Conversion {ext.upper()} to {fmt.upper()} is not supported.")
    if ext=="pdf":
        if fmt=="docx": return pdf_to_docx(src,outdir)
        if fmt=="xlsx": return pdf_to_xlsx(src,outdir)
        if fmt=="pptx": return pdf_to_pptx(src,outdir)
    if fmt=="pdf" and ext in ("docx","xlsx","pptx"): return libreoffice(src,outdir,"pdf")
    if ext in ("jpg","jpeg","png","webp") and fmt in ("jpg","jpeg","png","webp"): return image_convert(src,fmt,outdir)
    if fmt=="pdf" and ext in ("jpg","jpeg","png"):
        out=outdir/(src.stem+".pdf")
        with Image.open(src) as im: im.convert("RGB").save(out,"PDF",resolution=150)
        return out
    if ext=="txt" and fmt=="docx": return text_to_docx(src,outdir)
    if ext=="txt" and fmt=="pdf": return text_to_pdf(src,outdir)
    if ext=="csv" and fmt=="xlsx": return csv_to_xlsx(src,outdir)
    if ext=="xlsx" and fmt=="csv": return xlsx_to_csv(src,outdir)
    if ext in ("mp4","mkv","mov","avi","webm","wav","mp3"): return audio_convert(src,fmt,outdir,bitrate)
    raise RuntimeError("This conversion is not implemented.")
def merge_pdfs(paths,outdir):
    if len(paths)<2: raise RuntimeError("Select at least two PDF files to merge.")
    merger=PdfMerger()
    try:
        for p in paths:
            if p.suffix.lower()!=".pdf": raise RuntimeError("Only PDF files can be merged.")
            merger.append(str(p))
        out=outdir/"merged.pdf"; merger.write(str(out)); return out
    finally: merger.close()
def split_pdf(src,outdir):
    if src.suffix.lower()!=".pdf": raise RuntimeError("Select a PDF file.")
    reader=PdfReader(str(src))
    if len(reader.pages)<1: raise RuntimeError("The PDF has no pages.")
    results=[]
    for i,page in enumerate(reader.pages,1):
        writer=PdfWriter(); writer.add_page(page); out=outdir/f"page-{i:03d}.pdf"
        with open(out,"wb") as f: writer.write(f)
        results.append(out)
    return results
